import re
from pathlib import Path
from docx import Document
from docx.shared import Pt, RGBColor, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls, qn

WORKSPACE_DIR = Path(r"c:\xampp\php\www\keuangan-app")
MD_PATH = WORKSPACE_DIR / "docs" / "google-drive-setup.md"
DOCX_PATH = WORKSPACE_DIR / "docs" / "google-drive-setup.docx"

# Color constants
COLOR_PRIMARY = RGBColor(79, 70, 229)    # Indigo 600
COLOR_TEXT_HEADING = RGBColor(15, 23, 42) # Slate 900
COLOR_TEXT_BODY = RGBColor(30, 41, 59)   # Slate 800
COLOR_MUTED = RGBColor(100, 116, 139)    # Slate 500

def set_cell_shading(cell, color_hex):
    shading_elm = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{color_hex}"/>')
    cell._tc.get_or_add_tcPr().append(shading_elm)

def set_cell_left_border(cell, color_hex, size="24"):
    tcPr = cell._tc.get_or_add_tcPr()
    tcBorders = parse_xml(f'''
        <w:tcBorders xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">
            <w:top w:val="none"/>
            <w:left w:val="single" w:sz="{size}" w:space="0" w:color="{color_hex}"/>
            <w:bottom w:val="none"/>
            <w:right w:val="none"/>
        </w:tcBorders>
    ''')
    tcPr.append(tcBorders)

def set_cell_all_borders(cell, color_hex, size="4"):
    tcPr = cell._tc.get_or_add_tcPr()
    tcBorders = parse_xml(f'''
        <w:tcBorders xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">
            <w:top w:val="single" w:sz="{size}" w:space="0" w:color="{color_hex}"/>
            <w:left w:val="single" w:sz="{size}" w:space="0" w:color="{color_hex}"/>
            <w:bottom w:val="single" w:sz="{size}" w:space="0" w:color="{color_hex}"/>
            <w:right w:val="single" w:sz="{size}" w:space="0" w:color="{color_hex}"/>
        </w:tcBorders>
    ''')
    tcPr.append(tcBorders)

def add_inline_runs(paragraph, text):
    # We will tokenise by markdown inline elements:
    # `code` (inline code)
    # **bold** (bold)
    # [text](url) (links)
    
    # Simple regex scanner to find tokens sequentially
    pattern = re.compile(r'(\*\*.*?\*\*|`.*?`|\[.*?\]\(.*?\))')
    parts = pattern.split(text)
    
    for part in parts:
        if not part:
            continue
            
        if part.startswith('**') and part.endswith('**'):
            # Bold text
            run = paragraph.add_run(part[2:-2])
            run.bold = True
            run.font.color.rgb = COLOR_TEXT_HEADING
        elif part.startswith('`') and part.endswith('`'):
            # Inline code
            run = paragraph.add_run(part[1:-1])
            run.font.name = 'Courier New'
            run.font.size = Pt(9.5)
            run.font.color.rgb = RGBColor(199, 37, 78) # Dark crimson-like color for code
        elif part.startswith('[') and ']' in part and '(' in part and part.endswith(')'):
            # Link [text](url)
            match = re.match(r'\[(.*?)\]\((.*?)\)', part)
            if match:
                link_text, link_url = match.groups()
                run = paragraph.add_run(link_text)
                run.font.color.rgb = COLOR_PRIMARY
                run.underline = True
                # We also append the url in brackets for printed documents
                run_url = paragraph.add_run(f" ({link_url})")
                run_url.font.size = Pt(9)
                run_url.font.color.rgb = COLOR_MUTED
        else:
            # Regular text
            # Replace raw URLs with link runs
            url_pattern = re.compile(r'(https?://[^\s]+)')
            sub_parts = url_pattern.split(part)
            for sub_part in sub_parts:
                if url_pattern.match(sub_part):
                    run = paragraph.add_run(sub_part)
                    run.font.color.rgb = COLOR_PRIMARY
                    run.underline = True
                else:
                    run = paragraph.add_run(sub_part)

def main():
    if not MD_PATH.exists():
        print(f"Error: {MD_PATH} not found.")
        return
        
    md_content = MD_PATH.read_text(encoding="utf-8")
    lines = md_content.splitlines()
    
    doc = Document()
    
    # Document Margins (1 inch)
    for section in doc.sections:
        section.top_margin = Inches(1)
        section.bottom_margin = Inches(1)
        section.left_margin = Inches(1)
        section.right_margin = Inches(1)
        
    # Default Normal Style
    normal_style = doc.styles['Normal']
    normal_style.font.name = 'Segoe UI'
    normal_style.font.size = Pt(11)
    normal_style.font.color.rgb = COLOR_TEXT_BODY
    
    # Parse state variables
    in_code_block = False
    code_content = []
    
    in_table = False
    table_headers = []
    table_rows = []
    
    def process_paragraph_line(stripped_line):
        # Determine paragraph block styling
        if stripped_line.startswith("Catatan:"):
            # Indigo warning block
            tbl = doc.add_table(rows=1, cols=1)
            tbl.autofit = False
            tbl.columns[0].width = Inches(6.5)
            cell = tbl.cell(0, 0)
            set_cell_shading(cell, "EEF2F6") # Very soft slate-indigo shade
            set_cell_left_border(cell, "4F46E5", "24") # 3pt left border
            
            p = cell.paragraphs[0]
            p.paragraph_format.left_indent = Inches(0.15)
            p.paragraph_format.right_indent = Inches(0.15)
            p.paragraph_format.space_before = Pt(6)
            p.paragraph_format.space_after = Pt(6)
            
            add_inline_runs(p, stripped_line)
        else:
            p = doc.add_paragraph()
            p.paragraph_format.space_after = Pt(8)
            p.paragraph_format.line_spacing = 1.15
            add_inline_runs(p, stripped_line)

    def write_code_block():
        nonlocal code_content
        if not code_content:
            return
            
        tbl = doc.add_table(rows=1, cols=1)
        tbl.autofit = False
        tbl.columns[0].width = Inches(6.5)
        cell = tbl.cell(0, 0)
        
        # Style the code container
        set_cell_shading(cell, "F8FAFC") # slate 50 background
        set_cell_all_borders(cell, "CBD5E1", "4") # slate 300 thin borders
        
        p = cell.paragraphs[0]
        p.paragraph_format.left_indent = Inches(0.15)
        p.paragraph_format.right_indent = Inches(0.15)
        p.paragraph_format.space_before = Pt(8)
        p.paragraph_format.space_after = Pt(8)
        
        code_text = "\n".join(code_content)
        run = p.add_run(code_text)
        run.font.name = 'Courier New'
        run.font.size = Pt(9.5)
        run.font.color.rgb = RGBColor(51, 65, 85) # slate 700 text color
        
        code_content = []

    def write_table():
        nonlocal table_headers, table_rows
        if not table_headers and not table_rows:
            return
            
        num_cols = len(table_headers) if table_headers else (len(table_rows[0]) if table_rows else 1)
        tbl = doc.add_table(rows=0, cols=num_cols)
        tbl.style = 'Table Grid'
        
        # Write headers
        if table_headers:
            row = tbl.add_row()
            for idx, cell_text in enumerate(table_headers):
                cell = row.cells[idx]
                set_cell_shading(cell, "0F172A") # slate 900 headers
                p = cell.paragraphs[0]
                p.paragraph_format.space_before = Pt(4)
                p.paragraph_format.space_after = Pt(4)
                run = p.add_run(cell_text)
                run.bold = True
                run.font.color.rgb = RGBColor(255, 255, 255) # white text
                
        # Write rows
        for row_data in table_rows:
            row = tbl.add_row()
            for idx, cell_text in enumerate(row_data):
                if idx < len(row.cells):
                    cell = row.cells[idx]
                    p = cell.paragraphs[0]
                    p.paragraph_format.space_before = Pt(4)
                    p.paragraph_format.space_after = Pt(4)
                    add_inline_runs(p, cell_text)
                    
        doc.add_paragraph().paragraph_format.space_after = Pt(12)
        table_headers = []
        table_rows = []

    # Write document title (Cover-like heading)
    title_p = doc.add_paragraph()
    title_p.paragraph_format.space_before = Pt(12)
    title_p.paragraph_format.space_after = Pt(4)
    title_run = title_p.add_run("Setup Google Drive via Google Apps Script")
    title_run.font.name = 'Segoe UI'
    title_run.font.size = Pt(22)
    title_run.bold = True
    title_run.font.color.rgb = COLOR_TEXT_HEADING
    
    meta_p = doc.add_paragraph()
    meta_p.paragraph_format.space_after = Pt(24)
    meta_run = meta_p.add_run("Dokumen Panduan Setup • Aplikasi E-Keuangan • MAN 2 Surakarta")
    meta_run.font.size = Pt(9.5)
    meta_run.font.color.rgb = COLOR_MUTED
    meta_run.italic = True
    
    # Process lines
    for line in lines:
        stripped = line.strip()
        
        # Skip title line as we wrote it above
        if stripped == "# Setup Google Drive via Google Apps Script":
            continue
            
        # Code block handling
        if stripped.startswith("```"):
            if in_code_block:
                write_code_block()
                in_code_block = False
            else:
                write_table()
                in_code_block = True
            continue
            
        if in_code_block:
            code_content.append(line)
            continue
            
        # Table handling
        if stripped.startswith("|"):
            cells = [c.strip() for c in stripped.split("|")[1:-1]]
            if all(re.match(r'^:?-+:?$', c) for c in cells):
                in_table = True
                continue
            if not in_table:
                table_headers = cells
                in_table = True
            else:
                table_rows.append(cells)
            continue
        elif in_table:
            write_table()
            in_table = False
            
        # Header handling
        if stripped.startswith("#"):
            write_table()
            level = 0
            while level < len(stripped) and stripped[level] == '#':
                level += 1
            
            header_text = stripped[level:].strip()
            h_p = doc.add_paragraph()
            h_p.paragraph_format.space_before = Pt(18)
            h_p.paragraph_format.space_after = Pt(6)
            
            run = h_p.add_run(header_text)
            run.bold = True
            run.font.color.rgb = COLOR_TEXT_HEADING
            
            if level == 2:
                run.font.size = Pt(14)
                # Add horizontal line below H2
                p_border = doc.add_paragraph()
                p_border.paragraph_format.space_before = Pt(0)
                p_border.paragraph_format.space_after = Pt(6)
                pBdr = parse_xml(r'<w:pBdr xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"><w:bottom w:val="single" w:sz="6" w:space="1" w:color="CBD5E1"/></w:pBdr>')
                p_border._p.get_or_add_pPr().append(pBdr)
            elif level == 3:
                run.font.size = Pt(12)
                run.font.color.rgb = COLOR_PRIMARY
            else:
                run.font.size = Pt(16)
                
            continue
            
        # List handling
        ul_match = re.match(r'^[\-\*]\s+(.*)$', stripped)
        ol_match = re.match(r'^(\d+)\.\s+(.*)$', stripped)
        
        if ul_match:
            write_table()
            p = doc.add_paragraph(style='List Bullet')
            p.paragraph_format.space_after = Pt(4)
            p.paragraph_format.line_spacing = 1.15
            add_inline_runs(p, ul_match.group(1))
            continue
        elif ol_match:
            write_table()
            p = doc.add_paragraph(style='List Number')
            p.paragraph_format.space_after = Pt(4)
            p.paragraph_format.line_spacing = 1.15
            add_inline_runs(p, ol_match.group(2))
            continue
            
        # Empty Line
        if not stripped:
            write_table()
            continue
            
        # Normal Paragraph
        write_table()
        process_paragraph_line(stripped)
        
    # Flush remaining blocks
    write_code_block()
    write_table()
    
    # Save document
    doc.save(DOCX_PATH)
    print(f"Success: Generated beautiful Word document at {DOCX_PATH}")

if __name__ == "__main__":
    main()
